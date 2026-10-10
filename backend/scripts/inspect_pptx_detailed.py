import pptx

prs = pptx.Presentation(r"..\Campus Relay.pptx")

def dump_shape(shape, indent=0):
    ind = "  " * indent
    lines = []
    lines.append(f"{ind}[{shape.name}, type: {shape.shape_type}]")
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            txt = p.text.strip()
            if txt:
                lines.append(f"{ind}  T: {txt}")
    if shape.has_table:
        for r_idx, row in enumerate(shape.table.rows):
            row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
            lines.append(f"{ind}  Table R{r_idx}: {' | '.join(row_txt)}")
    if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            lines.extend(dump_shape(child, indent + 1))
    return lines

with open(r"..\docs\hackathon\current_pptx_detailed.txt", "w", encoding="utf-8") as f:
    f.write(f"Slide dimensions: {prs.slide_width} x {prs.slide_height} (aspect ratio: {prs.slide_width / prs.slide_height:.2f})\n")
    for idx, slide in enumerate(prs.slides, 1):
        f.write(f"\n==================== SLIDE {idx} ====================\n")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            f.write(f"[NOTES]: {slide.notes_slide.notes_text_frame.text.strip()}\n")
        for s_idx, shape in enumerate(slide.shapes):
            f.write(f"Top-level Shape {s_idx}:\n")
            for line in dump_shape(shape, 1):
                f.write(f"{line}\n")

print("Detailed dump completed.")
