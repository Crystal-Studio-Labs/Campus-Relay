import pptx

prs = pptx.Presentation(r"..\Campus Relay.pptx")
with open(r"..\docs\hackathon\current_pptx_dump.txt", "w", encoding="utf-8") as f:
    f.write(f"Slide dimensions: {prs.slide_width} x {prs.slide_height}\n")
    for idx, slide in enumerate(prs.slides, 1):
        f.write(f"\n==================== SLIDE {idx} ====================\n")
        # Check notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            f.write(f"[NOTES]: {slide.notes_slide.notes_text_frame.text.strip()}\n")
        for s_idx, shape in enumerate(slide.shapes):
            f.write(f"Shape {s_idx} [{shape.name}, type: {shape.shape_type}, left: {shape.left}, top: {shape.top}, width: {shape.width}, height: {shape.height}]:\n")
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    txt = p.text.strip()
                    if txt:
                        f.write(f"  P{p_idx}: {txt}\n")
            if shape.has_table:
                f.write("  [TABLE]:\n")
                for r_idx, row in enumerate(shape.table.rows):
                    row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                    f.write(f"    Row {r_idx}: {' | '.join(row_txt)}\n")
print("Dump completed.")
